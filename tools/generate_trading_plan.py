from __future__ import annotations

import re
import textwrap
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "trading-os-engineering-plan.md"
OUTPUT = ROOT / "docs" / "generated" / "Autonomous_Trading_OS_Engineering_Plan.docx"
DIAGRAMS = ROOT / "docs" / "diagrams"
RENDERED = DIAGRAMS / "rendered"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    run.font.size = Pt(8)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_toc(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Update this table in Word to populate page numbers."
    separate.append(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.05

    for name, size, color in (("Title", 25, "17365D"), ("Heading 1", 17, "17365D"), ("Heading 2", 13, "1F4E79"), ("Heading 3", 10.5, "2F5597")):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Cascadia Mono"
        code.font.size = Pt(8)
        code.font.color.rgb = RGBColor.from_string("243447")
        code.paragraph_format.left_indent = Inches(0.2)
        code.paragraph_format.right_indent = Inches(0.2)
        code.paragraph_format.space_before = Pt(3)
        code.paragraph_format.space_after = Pt(5)
    if "Callout" not in styles:
        callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
        callout.font.name = "Aptos"
        callout.font.size = Pt(9)
        callout.font.color.rgb = RGBColor.from_string("7F1D1D")
        callout.paragraph_format.left_indent = Inches(0.2)
        callout.paragraph_format.right_indent = Inches(0.2)
        callout.paragraph_format.space_before = Pt(4)
        callout.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "TSD · Autonomous Trading Operating System Engineering Plan · Proposed"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Aptos"
        run.font.size = Pt(7.5)
        run.font.color.rgb = RGBColor.from_string("64748B")
    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def parse_table_row(line: str) -> list[str]:
    body = line.strip().strip("|")
    return [cell.strip() for cell in body.split("|")]


def is_table_separator(line: str) -> bool:
    cells = parse_table_row(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells)


def add_inline(paragraph, text: str) -> None:
    # Lightweight Markdown inline formatting sufficient for this engineering document.
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]*\))")
    position = 0
    for match in token_re.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Cascadia Mono"
            run.font.size = Pt(8.5)
            run.font.color.rgb = RGBColor.from_string("334155")
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            label, url = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(label)
            run.underline = True
            run.font.color.rgb = RGBColor.from_string("0563C1")
            run._r.get_or_add_rPr().append(OxmlElement("w:hyperlink"))
        position = match.end()
    if position < len(text):
        paragraph.add_run(text[position:])


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=1, cols=columns)
    table.style = "Table Grid"
    table.autofit = True
    header = table.rows[0]
    set_repeat_table_header(header)
    for index in range(columns):
        cell = header.cells[index]
        cell.text = rows[0][index] if index < len(rows[0]) else ""
        set_cell_shading(cell, "D9EAF7")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(8)
    for row_values in rows[1:]:
        row = table.add_row()
        for index in range(columns):
            cell = row.cells[index]
            cell.text = row_values[index] if index < len(row_values) else ""
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(7.7)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_bullet(doc: Document, text: str, level: int = 0) -> None:
    paragraph = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(2)
    add_inline(paragraph, text)


def add_number(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(2)
    add_inline(paragraph, text)


def add_mermaid_snapshot(path: Path, output: Path) -> None:
    lines = path.read_text(encoding="utf-8").splitlines()
    title = path.stem.replace("-", " ").title()
    try:
        font = ImageFont.truetype("C:/Windows/Fonts/consola.ttf", 18)
        title_font = ImageFont.truetype("C:/Windows/Fonts/segoeui.ttf", 25)
    except OSError:
        font = ImageFont.load_default()
        title_font = font
    wrapped: list[str] = []
    for line in lines:
        wrapped.extend(textwrap.wrap(line, width=88, replace_whitespace=False) or [""])
    width = 1600
    height = max(160, 90 + len(wrapped) * 28)
    image = Image.new("RGB", (width, height), "#F8FAFC")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 0, width, 60), fill="#17365D")
    draw.text((28, 15), title, fill="white", font=title_font)
    y = 82
    for line in wrapped:
        draw.text((30, y), line, fill="#243447", font=font)
        y += 28
    output.parent.mkdir(parents=True, exist_ok=True)
    image.save(output, "PNG", optimize=True)


def add_diagram_appendix(doc: Document) -> None:
    doc.add_heading("Appendix A — Architecture diagram source snapshots", level=1)
    intro = doc.add_paragraph()
    add_inline(intro, "The source-controlled Mermaid files are the authoritative diagram definitions. PNG snapshots are embedded here so the DOCX remains readable without a Mermaid renderer.")
    for source in sorted(DIAGRAMS.glob("*.mmd")):
        image_path = RENDERED / f"{source.stem}.png"
        add_mermaid_snapshot(source, image_path)
        doc.add_heading(source.stem.replace("-", " ").title(), level=2)
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(image_path), width=Inches(6.6))
        caption = doc.add_paragraph(f"Figure A-{len(doc.inline_shapes)}: {source.name}")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption.runs[0].italic = True
        caption.runs[0].font.size = Pt(8)


def render_markdown(doc: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                paragraph = doc.add_paragraph(style="Code Block")
                paragraph.paragraph_format.left_indent = Inches(0.25)
                add_inline(paragraph, "\n".join(code_lines))
                in_code = False
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if not line.strip():
            index += 1
            continue
        if line.startswith("|") and index + 1 < len(lines) and lines[index + 1].startswith("|") and is_table_separator(lines[index + 1]):
            table_rows = [parse_table_row(line)]
            index += 2
            while index < len(lines) and lines[index].startswith("|"):
                table_rows.append(parse_table_row(lines[index]))
                index += 1
            add_table(doc, table_rows)
            continue
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2).strip()
            if level == 1:
                doc.add_heading(title, level=0)
            else:
                doc.add_heading(title, level=level - 1)
            index += 1
            continue
        if line.startswith(">"):
            paragraph = doc.add_paragraph(style="Callout")
            set_cell = paragraph._p.get_or_add_pPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "FEF2F2")
            set_cell.append(shading)
            add_inline(paragraph, line[1:].strip())
            index += 1
            continue
        bullet = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet:
            add_bullet(doc, bullet.group(2), 1 if len(bullet.group(1)) >= 2 else 0)
            index += 1
            continue
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if numbered:
            add_number(doc, numbered.group(1))
            index += 1
            continue
        if re.fullmatch(r"-{3,}", line.strip()):
            index += 1
            continue
        paragraph = doc.add_paragraph()
        add_inline(paragraph, line)
        index += 1


def add_cover_and_toc(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Autonomous Trading Operating System")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("TSD Engineering Plan and Foundation Decision")
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string("1F4E79")
    doc.add_paragraph()
    table = doc.add_table(rows=5, cols=2)
    table.style = "Light Shading Accent 1"
    metadata = [
        ("Status", "Proposed implementation baseline"),
        ("Version", "0.1"),
        ("Date", "2026-09-22"),
        ("Repository", "Multiagent_TSD-250926"),
        ("Safety status", "Live trading disabled until gates pass"),
    ]
    for row, (label, value) in zip(table.rows, metadata):
        row.cells[0].text = label
        row.cells[1].text = value
        row.cells[0].paragraphs[0].runs[0].bold = True
        for cell in row.cells:
            set_cell_margins(cell)
    doc.add_paragraph()
    warning = doc.add_paragraph(style="Callout")
    add_inline(warning, "This document is a design and delivery plan. Existing broker adapters and UI are prototypes; no live trading is authorized by this document.")
    doc.add_page_break()
    doc.add_heading("Table of contents", level=1)
    toc = doc.add_paragraph()
    add_toc(toc)
    doc.add_page_break()


def validate_output(path: Path) -> None:
    required = ["word/document.xml", "[Content_Types].xml", "_rels/.rels"]
    with ZipFile(path) as package:
        names = set(package.namelist())
        missing = [name for name in required if name not in names]
        if missing:
            raise RuntimeError(f"DOCX missing package parts: {missing}")
        if package.testzip() is not None:
            raise RuntimeError("DOCX ZIP integrity check failed")
    if path.stat().st_size < 50_000:
        raise RuntimeError("DOCX unexpectedly small; generation likely failed")


def main() -> None:
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover_and_toc(doc)
    render_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    add_diagram_appendix(doc)
    doc.save(OUTPUT)
    validate_output(OUTPUT)
    print(f"Generated {OUTPUT}")
    print(f"Size: {OUTPUT.stat().st_size:,} bytes")


if __name__ == "__main__":
    main()
